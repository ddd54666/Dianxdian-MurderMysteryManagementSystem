"""从 Excel 批量导入 DM 姓名，跳过数据库中已存在的同名记录。

使用方法：
    在项目根目录执行 python -m pip install openpyxl，随后执行 python import_dms.py。
    默认读取当前工作目录中的 DM.xlsx，写入项目根目录的 database.db。
    自定义文件路径时调用 import_dms(xlsx_path=文件路径)，主入口不解析命令行参数。
    导入前建议停止正在运行的系统并备份数据库；导入会直接修改现有 DM 档案。

Excel 格式：
    使用 .xlsx 文件，将目标工作表保存为活动工作表；脚本只读取 wb.active。
    第1行为表头，从第2行开始导入；至少保留 A-B 两列，不校验表头文字。
    A列：序号，可自行填写，脚本忽略其内容。
    B列：DM 姓名，必填文本；会去除首尾空格和单元格中的换行符。
    C列及之后的数据忽略；备注不会导入。请勿将姓名放在 A列或使用公式、合并单元格。

示例行：
    1 | 示例DM

处理规则：
    空姓名行跳过；同名 DM 只跳过，不更新原有记录，可重复执行。
    单行写入失败会输出错误并继续，有效记录在循环结束后提交，不保证整批同时成功。
"""

import os
import sys
import sqlite3
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook
from backend.config import get_db_path
from backend.models.database import _DDL_SCRIPT


def import_dms(xlsx_path='DM.xlsx'):
    """读取指定工作簿活动工作表的B列姓名，导入数据库中尚不存在的DM。"""
    if not os.path.exists(xlsx_path):
        print(f'[错误] 找不到文件: {xlsx_path}')
        return

    db_path = get_db_path()
    db = sqlite3.connect(db_path)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys = ON")

    # 确保表存在
    db.executescript(_DDL_SCRIPT)

    # 按原始单元格值读取，不计算公式；其他工作表不参与导入。
    wb = load_workbook(xlsx_path)
    ws = wb.active

    inserted = 0
    skipped = 0
    total = 0
    now = datetime.now().strftime('%Y-%m-%dT%H:%M:%S')

    # 首行作为表头跳过，姓名固定读取B列，A列序号不参与业务处理。
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 2:
            continue

        # 列: A=序号(忽略), B=DM名
        cell = row[1]
        if cell is None:
            continue
        name = str(cell).strip().replace('\n', '').replace('\r', '')
        if not name:
            continue

        total += 1

        # 检查重复
        existing = db.execute(
            'SELECT id FROM dms WHERE name = ?', (name,)
        ).fetchone()

        if existing:
            print(f'[跳过] DM已存在: {name}')
            skipped += 1
            continue

        try:
            db.execute(
                'INSERT INTO dms (name, created_at) VALUES (?, ?)',
                (name, now)
            )
            print(f'[导入] {name}')
            inserted += 1
        except Exception as e:
            print(f'[错误] 导入失败 {name}: {e}')

    db.commit()
    db.close()
    wb.close()
    print(f'\n完成: 导入 {inserted} 条, 跳过 {skipped} 条 (有效数据 {total} 行)')


if __name__ == '__main__':
    import_dms()
